"""Every kind of file she gets: Word, PowerPoint, old Excel, RTF, web pages, XML, and scans through OCR."""

import sys

import pytest
from helpers import write_pdf

from aagam import library, ocr
from aagam.db import open_db
from csdesk.documents.intake import extract_pages


@pytest.fixture
def db(settings):
    with open_db(settings.db_path, actor="test") as vault:
        yield vault


def test_powerpoint_one_page_per_slide_with_tables_and_notes(tmp_path):
    from pptx import Presentation
    from pptx.util import Inches

    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[1])
    slide.shapes.title.text = "Board presentation: Q2 FY 2026-27"
    slide.placeholders[1].text = "Borrowing limit under section 180(1)(c)"
    slide.notes_slide.notes_text_frame.text = "Seek approval for Rs. 100 crore"
    second = deck.slides.add_slide(deck.slide_layouts[5])
    table = second.shapes.add_table(2, 2, Inches(1), Inches(1), Inches(4), Inches(1)).table
    table.cell(0, 0).text, table.cell(0, 1).text, table.cell(1, 0).text, table.cell(1, 1).text = "Lender", "Limit", "SBI", "Rs. 60 crore"
    path = tmp_path / "board deck.pptx"
    deck.save(path)
    pages = extract_pages(path)
    assert len(pages) == 2
    assert "section 180(1)(c)" in pages[0].text and "Notes: Seek approval for Rs. 100 crore" in pages[0].text
    assert "SBI | Rs. 60 crore" in pages[1].text


def test_old_excel_files(tmp_path):
    xlwt = pytest.importorskip("xlwt")
    book = xlwt.Workbook()
    sheet = book.add_sheet("Trial balance")
    for r, row in enumerate([["Account", "Debit", "Credit"], ["Share capital", "", 50000000], ["Trade receivables", 61000000, ""]]):
        for c, value in enumerate(row):
            sheet.write(r, c, value)
    path = tmp_path / "TB FY26.xls"
    book.save(str(path))
    pages = extract_pages(path)
    assert pages[0].text.splitlines()[0] == "Sheet: Trial balance"
    assert "Trade receivables\t61000000\t" in pages[0].text  # whole numbers stay whole


def test_rtf_html_and_xml(tmp_path):
    rtf = tmp_path / "minutes.rtf"
    rtf.write_text(r"{\rtf1\ansi{\fonttbl\f0\fswiss Helvetica;}\f0\pard Minutes of the Board meeting held on 12th August, 2026.\par}", encoding="utf-8")
    assert "Minutes of the Board meeting held on 12th August, 2026." in extract_pages(rtf)[0].text
    html = tmp_path / "circular.html"
    html.write_text("<html><head><style>p{}</style></head><body><h1>General Circular No. 3/2026</h1><p>Extension of time</p>"
                    "<table><tr><td>Form</td><td>AOC-4</td></tr></table><script>x()</script></body></html>", encoding="utf-8")
    text = extract_pages(html)[0].text
    assert "General Circular No. 3/2026" in text and "Form | AOC-4" in text and "x()" not in text
    xml = tmp_path / "xbrl.xml"
    xml.write_text("<xbrl><in-gaap:RevenueFromOperations>620000000</in-gaap:RevenueFromOperations></xbrl>", encoding="utf-8")
    assert "RevenueFromOperations" in extract_pages(xml)[0].text


@pytest.mark.skipif(sys.platform != "win32" or not ocr.available(), reason="needs Windows OCR")
def test_a_scan_is_read_by_ocr_and_becomes_searchable(db, settings, tmp_path):
    from PIL import Image, ImageDraw, ImageFont

    font = ImageFont.truetype("arial.ttf", 34)
    image = Image.new("RGB", (1700, 900), "white")
    draw = ImageDraw.Draw(image)
    for i, line in enumerate(["CERTIFIED TRUE COPY OF THE RESOLUTION", "passed by the Board of Directors of",
                              "Acme Widgets Private Limited", "CIN: U74999MH2015PTC123456", "RESOLVED THAT the Company do open a",
                              "current account with State Bank of India."]):
        draw.text((80, 80 + i * 70), line, fill="black", font=font)
    path = tmp_path / "scanned ctc.png"
    image.save(path)
    info = library.add_document(db, settings, path, queue_ai=False)
    doc = library.document(db, info["document_id"])
    assert doc["ocr_pages"] == [1]
    assert doc["doc_type"] == "board_resolution"
    assert doc["facts"]["cin"][0]["value"] == "U74999MH2015PTC123456"
    assert library.search(db, "State Bank")["documents"][0]["document_id"] == info["document_id"]
    assert doc["scanned_pages"] == [1]  # still a scan, so Claude Desktop gets the image to read


def test_long_scans_wait_for_the_background(db, settings, tmp_path, monkeypatch):
    monkeypatch.setattr(ocr, "available", lambda: True)
    monkeypatch.setattr(ocr, "ocr_document", lambda vault, home, document_id, pages=None: {"read": []})
    info = library.add_document(db, settings, write_pdf(tmp_path / "scan.pdf", [""] * 8))
    jobs = {j["task"]: j["status"] for j in db._rows("SELECT task, status FROM ai_jobs WHERE document_id = ?", (info["document_id"],))}
    assert jobs == {"ocr": "queued"}  # the summary is queued only after OCR
