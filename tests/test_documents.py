import docx
import pytest
from helpers import ACME, USER, write_pdf

from csdesk.documents import ingest, ingest_inbox, read_pages
from csdesk.documents.intake import parse_page_spec
from csdesk.errors import Invalid
from csdesk.vault import Source

RESOLUTION = "CERTIFIED TRUE COPY OF THE RESOLUTION\nRESOLVED THAT Asha Rao (DIN 01234567) be appointed\nas an Additional Director"


def test_a_pdf_is_stored_indexed_and_its_scanned_page_flagged(vault, settings, tmp_path):
    company_id = vault.save("company", ACME, USER)
    pdf = write_pdf(tmp_path / "Board resolution.pdf", [RESOLUTION, ""])
    info = ingest(vault, settings, pdf, company_id, "board_resolution")
    assert info["pages"] == 2 and info["scanned_pages"] == [2]
    assert info["stored_path"].startswith(f"clients/{ACME['cin']}/documents/")
    assert (settings.home / info["stored_path"]).is_file()

    document, parts, rest = read_pages(vault, settings, info["document_id"], "all")
    assert document["original_name"] == "Board resolution.pdf"
    assert "DIN 01234567" in parts[0].text
    assert parts[1].text is None and parts[1].image.startswith(b"\xff\xd8\xff")  # the scan comes back as a JPEG
    assert rest == []


def test_the_same_file_is_not_stored_twice(vault, settings, tmp_path):
    pdf = write_pdf(tmp_path / "a.pdf", [RESOLUTION])
    first = ingest(vault, settings, pdf)
    again = ingest(vault, settings, pdf)
    assert again["already_on_file"] and again["document_id"] == first["document_id"]


def test_word_tables_keep_their_cells(vault, settings, tmp_path):
    path = tmp_path / "consent.docx"
    document = docx.Document()
    document.add_paragraph("Form DIR-2: consent to act as a director")
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text, table.rows[0].cells[1].text = "DIN", "01234567"
    document.save(str(path))
    info = ingest(vault, settings, path, doc_type="consent_dir2")
    _, parts, _ = read_pages(vault, settings, info["document_id"])
    assert "DIN | 01234567" in parts[0].text


def test_unsupported_files_are_refused(vault, settings, tmp_path):
    old = tmp_path / "minutes.doc"
    old.write_bytes(b"old word file")
    with pytest.raises(Invalid, match="can't read .doc"):
        ingest(vault, settings, old)


def test_the_inbox_is_emptied_into_processed(vault, settings):
    (settings.inbox_dir / "note.txt").write_text("Board meeting held on 26/09/2026", encoding="utf-8")
    (settings.inbox_dir / "old.doc").write_bytes(b"x")
    results = {r["file"]: r for r in ingest_inbox(vault, settings)}
    assert results["note.txt"]["document_id"]
    assert "can't read" in results["old.doc"]["error"]
    assert not (settings.inbox_dir / "note.txt").exists()
    assert (settings.inbox_dir / "old.doc").exists()  # left for her to deal with
    assert list((settings.inbox_dir / "processed").rglob("note.txt"))


def test_a_value_can_cite_a_page_that_exists_and_no_other(vault, settings, tmp_path):
    info = ingest(vault, settings, write_pdf(tmp_path / "coi.pdf", ["Certificate of Incorporation\nRoC-Mumbai"]))
    with pytest.raises(Invalid, match="page 2 doesn't exist"):
        vault.save("company", {"name": "X"}, Source("document", "test", document_id=info["document_id"], page=2))
    company_id = vault.save(
        "company", {"name": "X", "roc": "RoC-Mumbai"},
        Source("document", "test", document_id=info["document_id"], page=1, quote="RoC-Mumbai"),
    )
    source = vault.latest_sources("company", company_id)["roc"]
    assert (source["source_kind"], source["document_id"], source["page"], source["quote"]) == (
        "document", info["document_id"], 1, "RoC-Mumbai")


def test_page_lists():
    assert parse_page_spec("1-3,5,8-", 9) == [1, 2, 3, 5, 8, 9]
    assert parse_page_spec("all", 3) == [1, 2, 3]
    assert parse_page_spec("2-99", 3) == [2, 3]
    with pytest.raises(Invalid):
        parse_page_spec("first", 3)
