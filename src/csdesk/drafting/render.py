"""Drafts from her own Word templates.

Templates are .docx files in AAGAM_HOME/templates containing Jinja tags such
as {{ company.name }} or {{ extra.meeting_date }}. `company`, `directors`,
`kmp` and `auditor` come from the company file. `extra` carries whatever the
template needs that the file doesn't hold (meeting date, venue, signatory).
Two filters suit Indian documents: {{ x | ddmmyyyy }} and {{ x | inr }} (1,00,000).
"""

from __future__ import annotations

import re
import zipfile
from datetime import date
from pathlib import Path
from typing import Any

import jinja2
from docxtpl import DocxTemplate

from ..config import Settings
from ..errors import Invalid, NotFound
from ..vault import Vault
from ..vault.models import DIRECTOR_ROLES

SAMPLE = "sample-board-resolution-additional-director.docx"
KMP_ROLES = {"ceo", "cfo", "company_secretary", "manager", "managing_director", "whole_time_director"}


def ddmmyyyy(value: Any) -> str:
    if not value:
        return ""
    return date.fromisoformat(str(value)).strftime("%d-%m-%Y")


def inr(value: Any) -> str:
    """Indian digit grouping: 12345678 -> 1,23,45,678."""
    if value in (None, ""):
        return ""
    digits = str(int(value))
    sign, digits = ("-", digits[1:]) if digits.startswith("-") else ("", digits)
    if len(digits) <= 3:
        return sign + digits
    head, tail = digits[:-3], digits[-3:]
    groups = []
    while len(head) > 2:
        groups.insert(0, head[-2:])
        head = head[:-2]
    if head:
        groups.insert(0, head)
    return sign + ",".join(groups) + "," + tail


def ensure_sample(templates_dir: Path) -> Path:
    """Write the sample template once, so there is something to try before her own templates go in."""
    path = templates_dir / SAMPLE
    if path.exists():
        return path
    import docx

    templates_dir.mkdir(parents=True, exist_ok=True)
    document = docx.Document()
    document.add_paragraph(
        "CERTIFIED TRUE COPY OF THE RESOLUTION PASSED BY THE BOARD OF DIRECTORS OF {{ company.name | upper }} "
        "AT ITS MEETING HELD ON {{ extra.meeting_date | ddmmyyyy }} AT {{ extra.venue }}"
    )
    document.add_paragraph("CIN: {{ company.cin }}")
    document.add_paragraph("APPOINTMENT OF {{ extra.appointee_name | upper }} AS AN ADDITIONAL DIRECTOR")
    document.add_paragraph(
        '"RESOLVED THAT pursuant to Section 161(1) of the Companies Act, 2013, the rules made under it and the '
        "Articles of Association of the Company, {{ extra.appointee_name }} (DIN: {{ extra.appointee_din }}), who has "
        "given consent in Form DIR-2 and a declaration in Form DIR-8, be and is hereby appointed as an Additional "
        "Director of the Company with effect from {{ extra.meeting_date | ddmmyyyy }}, to hold office up to the date "
        "of the next Annual General Meeting or the last date on which it should have been held, whichever is earlier."
    )
    document.add_paragraph(
        "RESOLVED FURTHER THAT any Director or the Company Secretary of the Company be and is hereby severally "
        "authorised to file Form DIR-12 with the Registrar of Companies and to do all acts and things necessary to "
        'give effect to this resolution."'
    )
    document.add_paragraph("Certified true copy\nFor {{ company.name }}")
    document.add_paragraph(
        "{{ extra.signatory_name }}\n{{ extra.signatory_designation }}\nDIN / Membership No.: {{ extra.signatory_id }}"
    )
    document.add_paragraph("Place: {{ extra.place }}\nDate: {{ today }}")
    document.save(str(path))
    return path


def template_fields(path: Path) -> dict[str, list[str]]:
    """The `extra` keys a template expects, read from its text with Word's formatting tags stripped."""
    with zipfile.ZipFile(path) as archive:
        xml = " ".join(
            archive.read(name).decode("utf-8", "ignore")
            for name in archive.namelist()
            if re.fullmatch(r"word/(document|header\d*|footer\d*)\.xml", name)
        )
    text = re.sub(r"<[^>]+>", "", xml)
    extra = sorted(set(re.findall(r"extra\.(\w+)", text)))
    uses = sorted(set(re.findall(r"\{[{%]-?\s*(?:if\s+|for\s+\w+\s+in\s+)?(\w+)", text)) - {"extra", "if", "for"})
    return {"extra": extra, "uses": uses}


def list_templates(templates_dir: Path) -> list[dict[str, Any]]:
    return [
        {"template": path.name, **template_fields(path)}
        for path in sorted(templates_dir.glob("*.docx"))
        if not path.name.startswith("~$")
    ]


def render(
    vault: Vault, settings: Settings, company_id: int, template: str, extra: dict[str, Any] | None = None
) -> dict[str, Any]:
    name = template if template.lower().endswith(".docx") else f"{template}.docx"
    if Path(name).name != name:
        raise Invalid("Give the template's file name only, not a path.")
    path = settings.templates_dir / name
    if not path.is_file():
        known = ", ".join(t["template"] for t in list_templates(settings.templates_dir)) or "none yet"
        raise NotFound(f"There is no template {name}. Templates on file: {known}.")
    extra = extra or {}
    missing = [key for key in template_fields(path)["extra"] if key not in extra]
    if missing:
        raise Invalid(f"The template {name} also needs these extra values: {', '.join(missing)}.")

    file = vault.company_file(company_id)
    current = [p for p in file["people"] if not p["date_of_cessation"]]
    statutory = [a for a in file["auditors"] if a["audit_type"] == "statutory" and not a["ended_on"]]
    context = {
        "company": file["company"],
        "directors": [p for p in current if p["role"] in DIRECTOR_ROLES],
        "kmp": [p for p in current if p["role"] in KMP_ROLES],
        "auditor": statutory[0] if statutory else {},
        "today": date.today().strftime("%d-%m-%Y"),
        "extra": extra,
    }
    environment = jinja2.Environment(autoescape=True)
    environment.filters.update(ddmmyyyy=ddmmyyyy, inr=inr)
    document = DocxTemplate(str(path))
    document.render(context, environment)

    folder = settings.clients_dir / vault.company_folder(company_id) / "drafts"
    folder.mkdir(parents=True, exist_ok=True)
    out = folder / f"{date.today():%Y%m%d}-{path.stem}.docx"
    n = 2
    while out.exists():
        out = folder / f"{date.today():%Y%m%d}-{path.stem}-{n}.docx"
        n += 1
    document.save(str(out))
    vault.log("render draft", {"company_id": company_id, "template": name, "path": str(out)})
    return {"path": str(out), "template": name}
